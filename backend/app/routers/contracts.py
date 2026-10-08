import json
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy.orm import Session
from docx import Document

from app.database.connection import get_db
from app.models import Contract, Clause, Finding, Obligation

from app.services.document_service import extract_clauses
from app.services.risk_service import analyze_clauses
from app.services.obligation_service import extract_obligations
from app.services.version_service import generate_contract_v2
from app.services.diff_service import compare_clauses
from app.services.ai_service import analyze_contract_with_ai


router = APIRouter(
    prefix="/contracts",
    tags=["Contracts"]
)


# ============================================================
# LIST ALL CONTRACTS
# ============================================================

@router.get("")
def list_contracts(
    db: Session = Depends(get_db)
):

    contracts = (
        db.query(Contract)
        .order_by(
            Contract.contract_name.asc(),
            Contract.version_number.asc(),
            Contract.id.asc()
        )
        .all()
    )

    return {
        "contracts": [
            {
                "id": contract.id,
                "name": contract.contract_name,
                "file_name": contract.file_name,
                "file_type": contract.file_type,
                "version": contract.version_number,
                "parent_contract_id": contract.parent_contract_id
            }
            for contract in contracts
        ]
    }


UPLOAD_DIR = Path("uploads")
UPLOAD_DIR.mkdir(exist_ok=True)


# ============================================================
# UPLOAD CONTRACT
# ============================================================

@router.post("/upload")
async def upload_contract(
    file: UploadFile = File(...),
    contract_name: str = Form(""),
    db: Session = Depends(get_db)
):

    if not file.filename:
        raise HTTPException(
            status_code=400,
            detail="File name is required"
        )

    extension = Path(file.filename).suffix.lower()

    if extension not in [".pdf", ".docx"]:
        raise HTTPException(
            status_code=400,
            detail="Only PDF and DOCX files are supported"
        )

    file_path = UPLOAD_DIR / file.filename

    content = await file.read()

    with open(file_path, "wb") as buffer:
        buffer.write(content)

    clean_contract_name = (
            contract_name or Path(file.filename).stem
    ).strip()

    if not clean_contract_name:
        raise HTTPException(
            status_code=400,
            detail="Contract name is required"
        )

    # ========================================================
    # MANUAL UPLOADS ARE ALWAYS INDEPENDENT CONTRACTS
    # ========================================================

    version_number = 1
    parent_contract_id = None

    # ========================================================
    # CREATE CONTRACT RECORD
    # ========================================================

    contract = Contract(
        contract_name=clean_contract_name,
        file_name=file.filename,
        file_type=extension.replace(".", "").upper(),
        version_number=version_number,
        parent_contract_id=parent_contract_id
    )

    db.add(contract)
    db.commit()
    db.refresh(contract)

    try:
        extracted_clauses = extract_clauses(
            str(file_path)
        )

    except Exception as exc:
        db.delete(contract)
        db.commit()

        raise HTTPException(
            status_code=500,
            detail=f"Failed to extract contract clauses: {str(exc)}"
        )

    saved_clauses = []

    for item in extracted_clauses:

        clause = Clause(
            contract_id=contract.id,
            clause_number=item.get("clause_number"),
            clause_title=item.get("clause_title"),
            clause_text=item.get("clause_text"),
            page_number=item.get("page_number")
        )

        db.add(clause)
        db.flush()

        saved_clauses.append({
            "id": clause.id,
            "clause_number": clause.clause_number,
            "clause_title": clause.clause_title,
            "clause_text": clause.clause_text,
            "page_number": clause.page_number
        })

    db.commit()

    return {
        "message": "Contract uploaded successfully",

        "contract": {
            "id": contract.id,
            "name": contract.contract_name,
            "file_name": contract.file_name,
            "file_type": contract.file_type,
            "version": contract.version_number
        },

        "analysis": {
            "clause_count": len(saved_clauses),
            "clauses": saved_clauses
        }
    }


# ============================================================
# DELETE CONTRACT
# ============================================================

@router.delete("/{contract_id}")
def delete_contract(
    contract_id: int,
    db: Session = Depends(get_db)
):

    contract = (
        db.query(Contract)
        .filter(Contract.id == contract_id)
        .first()
    )

    if not contract:
        raise HTTPException(
            status_code=404,
            detail="Contract not found"
        )

    # --------------------------------------------------------
    # Find dependent generated versions
    # --------------------------------------------------------

    child_contracts = (
        db.query(Contract)
        .filter(
            Contract.parent_contract_id == contract_id
        )
        .all()
    )

    contracts_to_delete = [
        contract
    ] + child_contracts

    deleted_ids = []

    # --------------------------------------------------------
    # Delete physical uploaded/generated files
    # --------------------------------------------------------

    for item in contracts_to_delete:

        file_path = UPLOAD_DIR / item.file_name

        try:
            if file_path.exists():
                file_path.unlink()
        except Exception:
            pass

        deleted_ids.append(item.id)

    # --------------------------------------------------------
    # Delete database records
    # --------------------------------------------------------

    for item in contracts_to_delete:

        db.query(Obligation).filter(
            Obligation.contract_id == item.id
        ).delete(
            synchronize_session=False
        )

        db.query(Finding).filter(
            Finding.contract_id == item.id
        ).delete(
            synchronize_session=False
        )

        db.query(Clause).filter(
            Clause.contract_id == item.id
        ).delete(
            synchronize_session=False
        )

        db.delete(item)

    db.commit()

    return {
        "message": "Contract data deleted successfully",
        "deleted_contract_ids": deleted_ids
    }


# ============================================================
# ANALYZE CONTRACT AGAINST PLAYBOOK + AI ENGINE
# ============================================================

@router.post("/{contract_id}/analyze")
def analyze_contract(
    contract_id: int,
    db: Session = Depends(get_db)
):
    """
    Analyze a contract using two complementary engines:

    1. Deterministic playbook analysis:
       - AUTHORITATIVE for STANDARD/RISKY/MISSING/AMBIGUOUS
       - AUTHORITATIVE for severity, expected value, actual value,
         and clause traceability

    2. AI semantic analysis:
       - Enriches the result with key points, explanations and
         supporting-document suggestions
       - Does not override the playbook assessment
    """

    contract = (
        db.query(Contract)
        .filter(Contract.id == contract_id)
        .first()
    )

    if not contract:
        raise HTTPException(
            status_code=404,
            detail="Contract not found"
        )

    clauses = (
        db.query(Clause)
        .filter(Clause.contract_id == contract_id)
        .all()
    )

    if not clauses:
        raise HTTPException(
            status_code=404,
            detail="No clauses found for this contract"
        )

    clause_data = [
        {
            "id": clause.id,
            "clause_number": clause.clause_number,
            "clause_title": clause.clause_title,
            "clause_text": clause.clause_text,
            "page_number": clause.page_number
        }
        for clause in clauses
    ]

    # ========================================================
    # AUTHORITATIVE PLAYBOOK ANALYSIS
    # ========================================================

    deterministic_findings = analyze_clauses(clause_data)

    # ========================================================
    # AI ENRICHMENT
    # ========================================================

    ai_report = None
    ai_findings = []
    ai_error = None

    try:
        contract_path = UPLOAD_DIR / contract.file_name

        if not contract_path.exists():
            raise FileNotFoundError(
                f"Uploaded contract file not found: {contract_path}"
            )

        ai_report = analyze_contract_with_ai(
            str(contract_path)
        )

        ai_findings = ai_report.get("findings", []) or []

    except Exception as exc:
        ai_error = str(exc)

    # Index AI findings by rule so they can enrich the deterministic
    # finding without changing its authoritative classification.
    ai_by_rule = {}

    for item in ai_findings:
        rule_id = item.get("rule_id")

        if rule_id:
            ai_by_rule[rule_id] = item

    # ========================================================
    # MERGE: PLAYBOOK RESULT + AI EXPLANATION
    # ========================================================

    findings = []

    for item in deterministic_findings:

        rule_id = item.get("rule_id")
        ai_item = ai_by_rule.get(rule_id, {})

        # Keep deterministic values authoritative.
        finding = {
            "clause_id": item.get("clause_id"),
            "rule_id": rule_id,
            "category": item.get("category"),
            "status": item.get("status"),
            "severity": item.get("severity"),
            "evidence": item.get("evidence"),
            "expected": item.get("expected"),
            "actual": item.get("actual"),
            "reason": item.get("reason"),
            "recommended_action": item.get("recommended_action")
        }

        # AI can improve the explanation/action only when the
        # deterministic engine did not already provide one.
        if not finding["reason"]:
            finding["reason"] = (
                ai_item.get("explanation")
                or ai_item.get("reason")
            )

        if not finding["recommended_action"]:
            finding["recommended_action"] = (
                ai_item.get("suggested_action")
                or ai_item.get("recommended_action")
            )

        findings.append(finding)

    # ========================================================
    # SAVE AUTHORITATIVE FINDINGS
    # ========================================================

    db.query(Finding).filter(
        Finding.contract_id == contract_id
    ).delete(
        synchronize_session=False
    )

    saved_findings = []

    for item in findings:

        finding = Finding(
            contract_id=contract_id,
            clause_id=item.get("clause_id"),
            rule_id=item.get("rule_id"),
            category=item.get("category"),
            status=item.get("status"),
            severity=item.get("severity"),
            evidence=item.get("evidence"),
            expected=item.get("expected"),
            actual=item.get("actual"),
            reason=item.get("reason"),
            recommended_action=item.get("recommended_action")
        )

        db.add(finding)
        db.flush()

        saved_findings.append({
            "id": finding.id,
            "rule_id": finding.rule_id,
            "category": finding.category,
            "status": finding.status,
            "severity": finding.severity,
            "evidence": finding.evidence,
            "expected": finding.expected,
            "actual": finding.actual,
            "reason": finding.reason,
            "recommended_action": finding.recommended_action,
            "clause_id": finding.clause_id
        })

    db.commit()

    # ========================================================
    # AUTHORITATIVE SUMMARY
    # ========================================================

    high_risk = sum(
        1
        for item in saved_findings
        if item["severity"] == "HIGH"
    )

    medium_risk = sum(
        1
        for item in saved_findings
        if item["severity"] == "MEDIUM"
    )

    standard = sum(
        1
        for item in saved_findings
        if item["status"] == "STANDARD"
    )

    risky = sum(
        1
        for item in saved_findings
        if item["status"] == "RISKY"
    )

    missing = sum(
        1
        for item in saved_findings
        if item["status"] == "MISSING"
    )

    ambiguous = sum(
        1
        for item in saved_findings
        if item["status"] == "AMBIGUOUS"
    )

    authoritative_summary = {
        "clause_count": len(clauses),
        "finding_count": len(saved_findings),
        "high_risk": high_risk,
        "medium_risk": medium_risk,
        "standard": standard,
        "risky": risky,
        "missing": missing,
        "ambiguous": ambiguous
    }

    # ========================================================
    # AI RESPONSE
    # ========================================================

    ai_status = (
        "COMPLETED"
        if ai_report is not None
        else "FALLBACK_DETERMINISTIC"
    )

    ai_response = {
        "status": ai_status,
        "total_clauses": len(clauses),
        "total_key_points": 0,
        "key_points": [],
        "risk_summary": authoritative_summary,
        "ai_risk_summary": {},
        "supporting_document_suggestions": [],
        "total_supporting_document_suggestions": 0
    }

    if ai_report is not None:
        ai_response.update({
            # Use DB clause count so the API stays consistent with
            # the contract actually stored in the backend.
            "total_clauses": len(clauses),

            "total_key_points": ai_report.get(
                "total_key_points", 0
            ),

            "key_points": ai_report.get(
                "key_points", []
            ),

            # This is the authoritative summary shown by the
            # backend/dashboard.
            "risk_summary": authoritative_summary,

            # Preserve the raw AI interpretation separately so it
            # remains available for demo/debugging.
            "ai_risk_summary": ai_report.get(
                "risk_summary", {}
            ),

            "supporting_document_suggestions": ai_report.get(
                "supporting_document_suggestions", []
            ),

            "total_supporting_document_suggestions": ai_report.get(
                "total_supporting_document_suggestions", 0
            )
        })

    if ai_error:
        ai_response["error"] = ai_error

    return {
        "message": "Contract analyzed successfully",

        "contract": {
            "id": contract.id,
            "name": contract.contract_name,
            "version": contract.version_number
        },

        "analysis_engine": ai_response,

        "summary": authoritative_summary,

        "findings": saved_findings
    }


# ============================================================
# EXTRACT CONTRACT OBLIGATIONS
# ============================================================

@router.post("/{contract_id}/obligations")
def analyze_obligations(
    contract_id: int,
    db: Session = Depends(get_db)
):

    contract = (
        db.query(Contract)
        .filter(Contract.id == contract_id)
        .first()
    )

    if not contract:
        raise HTTPException(
            status_code=404,
            detail="Contract not found"
        )

    clauses = (
        db.query(Clause)
        .filter(
            Clause.contract_id == contract_id
        )
        .all()
    )

    if not clauses:
        raise HTTPException(
            status_code=404,
            detail="No clauses found for this contract"
        )

    clause_data = [
        {
            "id": clause.id,
            "clause_number": clause.clause_number,
            "clause_title": clause.clause_title,
            "clause_text": clause.clause_text,
            "page_number": clause.page_number
        }
        for clause in clauses
    ]

    obligations = extract_obligations(
        clause_data
    )

    db.query(Obligation).filter(
        Obligation.contract_id == contract_id
    ).delete(
        synchronize_session=False
    )

    saved_obligations = []

    for item in obligations:

        obligation = Obligation(
            contract_id=contract_id,
            clause_id=item.get("clause_id"),
            actor=item.get("actor"),
            action=item.get("action"),
            deadline=item.get("deadline"),
            trigger_condition=item.get(
                "trigger_condition"
            ),
            evidence=item.get("evidence")
        )

        db.add(obligation)
        db.flush()

        saved_obligations.append({
            "id": obligation.id,
            "clause_id": obligation.clause_id,
            "actor": obligation.actor,
            "action": obligation.action,
            "deadline": obligation.deadline,
            "trigger_condition": obligation.trigger_condition,
            "evidence": obligation.evidence
        })

    db.commit()

    return {
        "message": "Obligations extracted successfully",

        "contract": {
            "id": contract.id,
            "name": contract.contract_name,
            "version": contract.version_number
        },

        "summary": {
            "obligation_count": len(
                saved_obligations
            )
        },

        "obligations": saved_obligations
    }


# ============================================================
# UNIFIED CONTRACT INTELLIGENCE
# ============================================================

@router.get("/{contract_id}/intelligence")
def get_contract_intelligence(
    contract_id: int,
    db: Session = Depends(get_db)
):

    contract = (
        db.query(Contract)
        .filter(Contract.id == contract_id)
        .first()
    )

    if not contract:
        raise HTTPException(
            status_code=404,
            detail="Contract not found"
        )

    clauses = (
        db.query(Clause)
        .filter(
            Clause.contract_id == contract_id
        )
        .all()
    )

    findings = (
        db.query(Finding)
        .filter(
            Finding.contract_id == contract_id
        )
        .all()
    )

    obligations = (
        db.query(Obligation)
        .filter(
            Obligation.contract_id == contract_id
        )
        .all()
    )

    clause_response = []

    for clause in clauses:

        clause_response.append({
            "id": clause.id,
            "clause_number": clause.clause_number,
            "clause_title": clause.clause_title,
            "clause_text": clause.clause_text,
            "page_number": clause.page_number
        })

    finding_response = []

    for finding in findings:

        finding_response.append({
            "id": finding.id,
            "rule_id": finding.rule_id,
            "category": finding.category,
            "status": finding.status,
            "severity": finding.severity,
            "evidence": finding.evidence,
            "expected": finding.expected,
            "actual": finding.actual,
            "reason": finding.reason,
            "recommended_action": finding.recommended_action,
            "clause_id": finding.clause_id
        })

    obligation_response = []

    for obligation in obligations:

        obligation_response.append({
            "id": obligation.id,
            "clause_id": obligation.clause_id,
            "actor": obligation.actor,
            "action": obligation.action,
            "deadline": obligation.deadline,
            "trigger_condition": obligation.trigger_condition,
            "evidence": obligation.evidence
        })

    high_risk = sum(
        1
        for finding in findings
        if finding.severity == "HIGH"
    )

    medium_risk = sum(
        1
        for finding in findings
        if finding.severity == "MEDIUM"
    )

    risky = sum(
        1
        for finding in findings
        if finding.status == "RISKY"
    )

    standard = sum(
        1
        for finding in findings
        if finding.status == "STANDARD"
    )

    missing = sum(
        1
        for finding in findings
        if finding.status == "MISSING"
    )

    ambiguous = sum(
        1
        for finding in findings
        if finding.status == "AMBIGUOUS"
    )

    return {
        "contract": {
            "id": contract.id,
            "name": contract.contract_name,
            "file_name": contract.file_name,
            "file_type": contract.file_type,
            "version": contract.version_number
        },

        "summary": {
            "clause_count": len(clauses),
            "finding_count": len(findings),
            "obligation_count": len(obligations),
            "high_risk": high_risk,
            "medium_risk": medium_risk,
            "standard": standard,
            "risky": risky,
            "missing": missing,
            "ambiguous": ambiguous
        },

        "clauses": clause_response,

        "findings": finding_response,

        "obligations": obligation_response
    }


# ============================================================
# CONTRACT RISK DASHBOARD
# ============================================================

@router.get("/{contract_id}/dashboard")
def get_contract_dashboard(
    contract_id: int,
    db: Session = Depends(get_db)
):

    # --------------------------------------------------------
    # Find contract
    # --------------------------------------------------------

    contract = (
        db.query(Contract)
        .filter(Contract.id == contract_id)
        .first()
    )

    if not contract:
        raise HTTPException(
            status_code=404,
            detail="Contract not found"
        )

    # --------------------------------------------------------
    # Get findings
    # --------------------------------------------------------

    findings = (
        db.query(Finding)
        .filter(Finding.contract_id == contract_id)
        .all()
    )

    # --------------------------------------------------------
    # Get obligations
    # --------------------------------------------------------

    obligations = (
        db.query(Obligation)
        .filter(Obligation.contract_id == contract_id)
        .all()
    )

    # --------------------------------------------------------
    # Risk summary
    # --------------------------------------------------------

    high_risk = sum(
        1
        for finding in findings
        if finding.severity == "HIGH"
    )

    medium_risk = sum(
        1
        for finding in findings
        if finding.severity == "MEDIUM"
    )

    standard = sum(
        1
        for finding in findings
        if finding.status == "STANDARD"
    )

    risky = sum(
        1
        for finding in findings
        if finding.status == "RISKY"
    )

    ambiguous = sum(
        1
        for finding in findings
        if finding.status == "AMBIGUOUS"
    )

    missing = sum(
        1
        for finding in findings
        if finding.status == "MISSING"
    )

    # --------------------------------------------------------
    # Risk ranking
    # --------------------------------------------------------

    severity_rank = {
        "HIGH": 3,
        "MEDIUM": 2,
        "LOW": 1
    }

    status_rank = {
        "RISKY": 3,
        "AMBIGUOUS": 2,
        "MISSING": 2,
        "STANDARD": 1
    }

    def risk_score(finding):
        return max(
            severity_rank.get(finding.severity, 0),
            status_rank.get(finding.status, 0)
        )

    sorted_findings = sorted(
        findings,
        key=risk_score,
        reverse=True
    )

    # --------------------------------------------------------
    # Top risks
    # --------------------------------------------------------

    top_risks = []

    for finding in sorted_findings:

        if finding.status not in [
            "RISKY",
            "AMBIGUOUS",
            "MISSING"
        ]:
            continue

        top_risks.append({
            "finding_id": finding.id,
            "rule_id": finding.rule_id,
            "category": finding.category,
            "status": finding.status,
            "severity": finding.severity,
            "actual": finding.actual,
            "expected": finding.expected,
            "evidence": finding.evidence,
            "reason": finding.reason,
            "recommended_action": finding.recommended_action,
            "clause_id": finding.clause_id,
            "trace_endpoint": (
                f"/contracts/{contract_id}/"
                f"findings/{finding.id}/trace"
            )
        })

    # --------------------------------------------------------
    # Missing clauses
    # --------------------------------------------------------

    missing_clauses = [
        {
            "finding_id": finding.id,
            "rule_id": finding.rule_id,
            "category": finding.category,
            "severity": finding.severity,
            "expected": finding.expected,
            "recommended_action": finding.recommended_action
        }
        for finding in findings
        if finding.status == "MISSING"
    ]

    # --------------------------------------------------------
    # Ambiguous clauses
    # --------------------------------------------------------

    ambiguous_clauses = [
        {
            "finding_id": finding.id,
            "rule_id": finding.rule_id,
            "category": finding.category,
            "severity": finding.severity,
            "evidence": finding.evidence,
            "actual": finding.actual,
            "expected": finding.expected,
            "reason": finding.reason,
            "recommended_action": finding.recommended_action,
            "clause_id": finding.clause_id
        }
        for finding in findings
        if finding.status == "AMBIGUOUS"
    ]

    # --------------------------------------------------------
    # Category distribution
    # --------------------------------------------------------

    category_distribution = {}

    for finding in findings:

        category = finding.category

        if category not in category_distribution:
            category_distribution[category] = {
                "category": category,
                "status": finding.status,
                "severity": finding.severity,
                "count": 0
            }

        category_distribution[category]["count"] += 1

    # --------------------------------------------------------
    # Obligations
    # --------------------------------------------------------

    obligation_response = [
        {
            "id": obligation.id,
            "clause_id": obligation.clause_id,
            "actor": obligation.actor,
            "action": obligation.action,
            "deadline": obligation.deadline,
            "trigger_condition": obligation.trigger_condition,
            "evidence": obligation.evidence
        }
        for obligation in obligations
    ]

    # --------------------------------------------------------
    # Internal risk indicator
    # --------------------------------------------------------

    risk_points = (
        (high_risk * 3)
        + (medium_risk * 2)
    )

    if risk_points >= 6:
        risk_level = "HIGH"
    elif risk_points >= 3:
        risk_level = "MEDIUM"
    else:
        risk_level = "LOW"

    # --------------------------------------------------------
    # Final dashboard response
    # --------------------------------------------------------

    return {
        "message": "Contract dashboard generated successfully",

        "contract": {
            "id": contract.id,
            "name": contract.contract_name,
            "file_name": contract.file_name,
            "file_type": contract.file_type,
            "version": contract.version_number
        },

        "risk_summary": {
            "total_findings": len(findings),
            "high_risk": high_risk,
            "medium_risk": medium_risk,
            "risky": risky,
            "standard": standard,
            "ambiguous": ambiguous,
            "missing": missing
        },

        "risk_indicator": {
            "level": risk_level,
            "points": risk_points,
            "disclaimer": (
                "Internal contract analysis indicator only; "
                "not legal advice."
            )
        },

        "top_risks": top_risks,

        "missing_clauses": missing_clauses,

        "ambiguous_clauses": ambiguous_clauses,

        "category_distribution": list(
            category_distribution.values()
        ),

        "obligations": obligation_response
    }


# ============================================================
# GENERATE CONTRACT V2
# ============================================================

@router.post("/{contract_id}/generate-v2")
def generate_v2(
    contract_id: int,
    db: Session = Depends(get_db)
):

    original_contract = (
        db.query(Contract)
        .filter(
            Contract.id == contract_id
        )
        .first()
    )

    if not original_contract:
        raise HTTPException(
            status_code=404,
            detail="Contract not found"
        )

    if original_contract.version_number != 1:
        raise HTTPException(
            status_code=400,
            detail="V2 generation must start from a Version 1 contract"
        )

    clauses = (
        db.query(Clause)
        .filter(
            Clause.contract_id == contract_id
        )
        .order_by(Clause.id)
        .all()
    )

    if not clauses:
        raise HTTPException(
            status_code=404,
            detail="No clauses found for this contract"
        )

    revised_clauses = generate_contract_v2(
        clauses
    )

    changed_count = sum(
        1
        for clause in revised_clauses
        if clause["changed"]
    )

    version_2 = Contract(
        contract_name=original_contract.contract_name,
        file_name=(
            f"{original_contract.contract_name}_v2.docx"
        ),
        file_type="DOCX",
        version_number=2,
        parent_contract_id=original_contract.id
    )

    db.add(version_2)
    db.commit()
    db.refresh(version_2)

    saved_clauses = []

    for item in revised_clauses:

        clause = Clause(
            contract_id=version_2.id,
            clause_number=item["clause_number"],
            clause_title=item["clause_title"],
            clause_text=item["clause_text"],
            page_number=item["page_number"]
        )

        db.add(clause)
        db.flush()

        saved_clauses.append({
            "id": clause.id,
            "clause_number": clause.clause_number,
            "clause_title": clause.clause_title,
            "clause_text": clause.clause_text,
            "changed": item["changed"]
        })

    db.commit()

    # ========================================================
    # CREATE PHYSICAL V2 DOCX FOR AI ANALYSIS
    # ========================================================

    v2_file_path = UPLOAD_DIR / version_2.file_name

    try:
        document = Document()

        for item in revised_clauses:
            clause_number = item.get("clause_number")
            clause_title = item.get("clause_title")
            clause_text = item.get("clause_text")

            if clause_number and clause_title:
                document.add_heading(
                    f"{clause_number}. {clause_title}",
                    level=2
                )
            elif clause_title:
                document.add_heading(
                    clause_title,
                    level=2
                )

            if clause_text:
                document.add_paragraph(clause_text)

        document.save(v2_file_path)
        # ========================================================
        # AUTOMATICALLY ANALYZE GENERATED V2
        # ========================================================

        try:
            analyze_contract(
                version_2.id,
                db
            )

            analyze_obligations(
                version_2.id,
                db
            )

        except Exception as exc:
            raise HTTPException(
                status_code=500,
                detail=(
                    "V2 was generated successfully, "
                    f"but automatic analysis failed: {str(exc)}"
                )
            )

    except Exception as exc:
        raise HTTPException(

            status_code=500,
            detail=f"Failed to create V2 DOCX: {str(exc)}"
        )

    return {
        "message": "Contract V2 generated successfully",

        "original_contract": {
            "id": original_contract.id,
            "version": original_contract.version_number
        },

        "new_contract": {
            "id": version_2.id,
            "version": version_2.version_number,
            "parent_contract_id": version_2.parent_contract_id,
            "name": version_2.contract_name
        },

        "summary": {
            "original_clause_count": len(clauses),
            "changed_clause_count": changed_count,
            "unchanged_clause_count": (
                len(clauses) - changed_count
            )
        },

        "clauses": saved_clauses
    }


# ============================================================
# COMPARE CONTRACT V1 WITH V2
# ============================================================

@router.get("/{contract_id}/compare-v2")
def compare_contract_v2(
    contract_id: int,
    db: Session = Depends(get_db)
):

    original_contract = (
        db.query(Contract)
        .filter(
            Contract.id == contract_id
        )
        .first()
    )

    if not original_contract:
        raise HTTPException(
            status_code=404,
            detail="Original contract not found"
        )

    version_2 = (
        db.query(Contract)
        .filter(
            Contract.parent_contract_id == original_contract.id,
            Contract.version_number == 2
        )
        .order_by(Contract.id.desc())
        .first()
    )

    if not version_2:
        raise HTTPException(
            status_code=404,
            detail="Version 2 not found. Generate V2 first."
        )

    v1_clauses = (
        db.query(Clause)
        .filter(
            Clause.contract_id == original_contract.id
        )
        .order_by(Clause.id)
        .all()
    )

    v2_clauses = (
        db.query(Clause)
        .filter(
            Clause.contract_id == version_2.id
        )
        .order_by(Clause.id)
        .all()
    )

    changes = compare_clauses(
        v1_clauses,
        v2_clauses
    )

    modified = sum(
        1
        for item in changes
        if item["change_type"] == "MODIFIED"
    )

    added = sum(
        1
        for item in changes
        if item["change_type"] == "ADDED"
    )

    removed = sum(
        1
        for item in changes
        if item["change_type"] == "REMOVED"
    )

    unchanged = sum(
        1
        for item in changes
        if item["change_type"] == "UNCHANGED"
    )

    return {
        "message": "Contract versions compared successfully",

        "comparison": {
            "original_contract": {
                "id": original_contract.id,
                "version": original_contract.version_number
            },

            "new_contract": {
                "id": version_2.id,
                "version": version_2.version_number
            }
        },

        "summary": {
            "total_clauses": len(changes),
            "modified": modified,
            "added": added,
            "removed": removed,
            "unchanged": unchanged
        },

        "changes": changes
    }


# ============================================================
# EVIDENCE → PLAYBOOK TRACEABILITY
# ============================================================

@router.get("/{contract_id}/findings/{finding_id}/trace")
def trace_finding(
    contract_id: int,
    finding_id: int,
    db: Session = Depends(get_db)
):

    contract = (
        db.query(Contract)
        .filter(Contract.id == contract_id)
        .first()
    )

    if not contract:
        raise HTTPException(
            status_code=404,
            detail="Contract not found"
        )

    finding = (
        db.query(Finding)
        .filter(
            Finding.id == finding_id,
            Finding.contract_id == contract_id
        )
        .first()
    )

    if not finding:
        raise HTTPException(
            status_code=404,
            detail="Finding not found for this contract"
        )

    clause = None

    if finding.clause_id:
        clause = (
            db.query(Clause)
            .filter(
                Clause.id == finding.clause_id
            )
            .first()
        )

    rules_path = (
        Path(__file__).resolve().parents[3]
        / "playbook"
        / "rules.json"
    )

    try:
        with open(
            rules_path,
            "r",
            encoding="utf-8"
        ) as file:
            playbook = json.load(file)

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Failed to load playbook rules: {str(exc)}"
        )

    rule = next(
        (
            item
            for item in playbook.get("rules", [])
            if item.get("rule_id") == finding.rule_id
        ),
        None
    )

    if not rule:
        raise HTTPException(
            status_code=404,
            detail=f"Playbook rule not found: {finding.rule_id}"
        )

    return {
        "message": "Finding trace generated successfully",

        "contract": {
            "id": contract.id,
            "name": contract.contract_name,
            "version": contract.version_number,
            "file_name": contract.file_name
        },

        "finding": {
            "id": finding.id,
            "rule_id": finding.rule_id,
            "category": finding.category,
            "status": finding.status,
            "severity": finding.severity
        },

        "contract_evidence": {
            "clause_id": (
                clause.id
                if clause
                else finding.clause_id
            ),

            "clause_number": (
                clause.clause_number
                if clause
                else None
            ),

            "clause_title": (
                clause.clause_title
                if clause
                else None
            ),

            "page_number": (
                clause.page_number
                if clause
                else None
            ),

            "text": (
                clause.clause_text
                if clause
                else finding.evidence
            )
        },

        "playbook_rule": {
            "rule_id": rule.get("rule_id"),
            "category": rule.get("category"),
            "requirement": rule.get("requirement"),
            "severity": rule.get("severity"),
            "description": rule.get("description")
        },

        "assessment": {
            "actual": finding.actual,
            "expected": finding.expected,
            "evidence": finding.evidence,
            "reason": finding.reason,
            "recommended_action": finding.recommended_action
        }
    }


# ============================================================
# COMPARE RISK BETWEEN V1 AND V2
# ============================================================

@router.get("/{contract_id}/risk-comparison")
def compare_risk_versions(
    contract_id: int,
    v2_contract_id: Optional[int] = None,
    db: Session = Depends(get_db)
):

    # The comparison UI always sends both IDs explicitly. There is
    # deliberately NO fallback: no parent lookup, no "latest V2", no
    # name/filename/family matching. (The parameter is Optional in the
    # signature only so a missing value returns 400 instead of
    # FastAPI's automatic 422.)

    if v2_contract_id is None:
        raise HTTPException(
            status_code=400,
            detail="v2_contract_id is required. Select the contract to compare."
        )

    if v2_contract_id == contract_id:
        raise HTTPException(
            status_code=400,
            detail="Select two different contracts to compare."
        )

    v1 = (
        db.query(Contract)
        .filter(Contract.id == contract_id)
        .first()
    )

    if not v1:
        raise HTTPException(
            status_code=404,
            detail="Original contract not found"
        )

    v2 = (
        db.query(Contract)
        .filter(Contract.id == v2_contract_id)
        .first()
    )

    if not v2:
        raise HTTPException(
            status_code=404,
            detail="Comparison contract not found"
        )

    v1_findings = (
        db.query(Finding)
        .filter(Finding.contract_id == v1.id)
        .all()
    )

    v2_findings = (
        db.query(Finding)
        .filter(Finding.contract_id == v2.id)
        .all()
    )

    if not v1_findings:
        raise HTTPException(
            status_code=400,
            detail="Version 1 has not been analyzed yet. Run /analyze first."
        )

    if not v2_findings:
        raise HTTPException(
            status_code=400,
            detail="Version 2 has not been analyzed yet. Run /analyze for V2 first."
        )

    v1_map = {
        finding.rule_id: finding
        for finding in v1_findings
    }

    v2_map = {
        finding.rule_id: finding
        for finding in v2_findings
    }

    all_rules = sorted(
        set(v1_map.keys()) | set(v2_map.keys())
    )

    severity_rank = {
        "HIGH": 3,
        "MEDIUM": 2,
        "LOW": 1
    }

    status_rank = {
        "RISKY": 3,
        "AMBIGUOUS": 2,
        "MISSING": 2,
        "STANDARD": 1
    }

    def risk_score(finding):
        return max(
            severity_rank.get(finding.severity, 0),
            status_rank.get(finding.status, 0)
        )

    def finding_payload(finding):
        return {
            "status": finding.status,
            "severity": finding.severity,
            "actual": finding.actual,
            "expected": finding.expected,
            "evidence": finding.evidence,
            "reason": finding.reason,
            "recommended_action": finding.recommended_action,
            "clause_id": finding.clause_id
        }

    risk_changes = []

    for rule_id in all_rules:

        old = v1_map.get(rule_id)
        new = v2_map.get(rule_id)

        if old and new:

            old_score = risk_score(old)
            new_score = risk_score(new)

            if new_score < old_score:
                impact = "RISK REDUCED"
            elif new_score > old_score:
                impact = "RISK INCREASED"
            else:
                impact = "NO CHANGE"

            risk_changes.append({
                "rule_id": rule_id,
                "category": old.category,
                "v1": finding_payload(old),
                "v2": finding_payload(new),
                "v1_risk_score": old_score,
                "v2_risk_score": new_score,
                "risk_impact": impact
            })

        elif old and not new:

            old_score = risk_score(old)

            if old_score >= 2:
                impact = "RISK RESOLVED"
            else:
                impact = "FINDING REMOVED"

            risk_changes.append({
                "rule_id": rule_id,
                "category": old.category,
                "v1": finding_payload(old),
                "v2": None,
                "v1_risk_score": old_score,
                "v2_risk_score": 0,
                "risk_impact": impact
            })

        elif new and not old:

            new_score = risk_score(new)

            if new_score >= 2:
                impact = "NEW RISK"
            else:
                impact = "NEW STANDARD FINDING"

            risk_changes.append({
                "rule_id": rule_id,
                "category": new.category,
                "v1": None,
                "v2": finding_payload(new),
                "v1_risk_score": 0,
                "v2_risk_score": new_score,
                "risk_impact": impact
            })

    risks_reduced = sum(
        1
        for item in risk_changes
        if item["risk_impact"] == "RISK REDUCED"
    )

    risks_increased = sum(
        1
        for item in risk_changes
        if item["risk_impact"] == "RISK INCREASED"
    )

    risks_resolved = sum(
        1
        for item in risk_changes
        if item["risk_impact"] == "RISK RESOLVED"
    )

    new_risks = sum(
        1
        for item in risk_changes
        if item["risk_impact"] == "NEW RISK"
    )

    unchanged = sum(
        1
        for item in risk_changes
        if item["risk_impact"] == "NO CHANGE"
    )

    risk_regression = (
        risks_increased > 0
        or new_risks > 0
    )

    net_risk_improvement = (
        risks_reduced
        + risks_resolved
        - risks_increased
        - new_risks
    )

    v1_high = sum(
        1
        for finding in v1_findings
        if finding.severity == "HIGH"
    )

    v2_high = sum(
        1
        for finding in v2_findings
        if finding.severity == "HIGH"
    )

    v1_risky = sum(
        1
        for finding in v1_findings
        if finding.status == "RISKY"
    )

    v2_risky = sum(
        1
        for finding in v2_findings
        if finding.status == "RISKY"
    )

    v1_medium = sum(
        1
        for finding in v1_findings
        if finding.severity == "MEDIUM"
    )

    v2_medium = sum(
        1
        for finding in v2_findings
        if finding.severity == "MEDIUM"
    )

    v1_missing = sum(
        1
        for finding in v1_findings
        if finding.status == "MISSING"
    )

    v2_missing = sum(
        1
        for finding in v2_findings
        if finding.status == "MISSING"
    )

    v1_ambiguous = sum(
        1
        for finding in v1_findings
        if finding.status == "AMBIGUOUS"
    )

    v2_ambiguous = sum(
        1
        for finding in v2_findings
        if finding.status == "AMBIGUOUS"
    )

    return {
        "message": "Risk comparison completed successfully",

        "versions": {
            "v1": {
                "id": v1.id,
                "version": v1.version_number,
                "name": v1.contract_name,
                "file_name": v1.file_name
            },

            "v2": {
                "id": v2.id,
                "version": v2.version_number,
                "name": v2.contract_name,
                "file_name": v2.file_name,
                "parent_contract_id": v2.parent_contract_id
            }
        },

        "risk_summary": {
            "v1_high_risk": v1_high,
            "v2_high_risk": v2_high,
            "v1_risky": v1_risky,
            "v2_risky": v2_risky,
            "v1_medium_risk": v1_medium,
            "v2_medium_risk": v2_medium,
            "v1_missing": v1_missing,
            "v2_missing": v2_missing,
            "v1_ambiguous": v1_ambiguous,
            "v2_ambiguous": v2_ambiguous,
            "high_risk_reduction": v1_high - v2_high,
            "risky_finding_reduction": v1_risky - v2_risky,
            "risks_reduced": risks_reduced,
            "risks_resolved": risks_resolved,
            "risks_increased": risks_increased,
            "new_risks": new_risks,
            "unchanged": unchanged,
            "net_risk_improvement": net_risk_improvement,
            "risk_regression": risk_regression
        },

        "risk_changes": risk_changes
    }
